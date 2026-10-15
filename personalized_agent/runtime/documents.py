"""READING WHAT IS NOT TEXT - pictures, PDFs, Word, PowerPoint and Excel - for
every model, whether it can see pictures or not.

  - A PICTURE: its words read here, on this computer, by the system's own
    OCR (platform_layer/ocr.py: Windows.Media.Ocr, macOS' Vision, RapidOCR
    or Tesseract on Linux - no picture leaves it for this); and the picture
    itself, made small enough,
    for a model that can see one (models/vision.py).
  - A PDF: the text of each page; a page that has none (a scan) has its
    words read off it the same way.
  - WORD, POWERPOINT, EXCEL: their text, and a workbook's cells.

ONE READING PER FILE: what was read of a file is kept as a COPY, named by
the SHA-256 of the file's bytes, in the application's own folder - never
beside the file. Read again, the copy is used and nothing is read twice; a
file whose bytes changed is a different file, and is read anew. The copies
least lately used go first when there are too many.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import logging
import os
import tempfile
import threading
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

log = logging.getLogger(__name__)

PICTURES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff")
DOCUMENTS = (".pdf", ".docx", ".pptx", ".xlsx", ".xlsm")
#: How a reading is kept: changed, every copy is read again.
VERSION = 1
#: Copies kept at most; the least lately used go first.
MOST_COPIES = 5000
#: Pages of one PDF whose words are read off their picture, at most.
MOST_OCR_PAGES = 40
#: A page with less text than this is taken for a scan.
SCANNED_BELOW = 30
#: The longest side of a picture given to a model, and to the OCR.
PICTURE_SIDE = 1568
OCR_SIDE = 2600
#: The most of a workbook read.
MOST_CELLS = 20000

_lock = threading.Lock()
_folder: Optional[Path] = None
#: Whether accounts and passwords are hidden from the words read off a
#: picture (Settings > Privacy) - asked each time.
_hiding: Any = lambda: True


def configure(folder: Optional[Path], hiding: Any = None) -> None:
    """Where the copies are kept: the application's own folder; and whether
    accounts and passwords are hidden from what is read off pictures."""
    global _folder, _hiding
    _folder = Path(folder) if folder else None
    if hiding is not None:
        _hiding = hiding
    if _folder is not None:
        _folder.mkdir(parents=True, exist_ok=True)


def kind_of(path: Any) -> str:
    suffix = Path(str(path)).suffix.lower()
    if suffix in PICTURES:
        return "picture"
    if suffix in DOCUMENTS:
        return {".xlsm": "xlsx"}.get(suffix, suffix[1:])
    return "text"


def needs_reading(path: Any) -> bool:
    """Whether a file is one this module reads (not plain text)."""
    return kind_of(path) != "text"


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


# -- the copies ------------------------------------------------------------------------------------
def _copy_path(sha: str) -> Optional[Path]:
    if _folder is None:
        return None
    return _folder / sha[:2] / f"{sha}.json"


def copy_of(path: Any) -> Optional[Dict[str, Any]]:
    """The reading kept of a file, if there is one - without reading it."""
    try:
        sha = sha256_of(Path(path))
    except OSError:
        return None
    return _kept(sha)


def _kept(sha: str) -> Optional[Dict[str, Any]]:
    where = _copy_path(sha)
    if where is None or not where.exists():
        return None
    try:
        said = json.loads(where.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if said.get("version") != VERSION:
        return None
    try:
        os.utime(where)                       # used: kept the longer
    except OSError:
        pass
    return said


def _keep(sha: str, said: Dict[str, Any]) -> None:
    where = _copy_path(sha)
    if where is None:
        return
    with _lock:
        where.parent.mkdir(parents=True, exist_ok=True)
        where.write_text(json.dumps(said, ensure_ascii=False), encoding="utf-8")
        _trim()


def _trim() -> None:
    copies = list(_folder.glob("*/*.json")) if _folder is not None else []
    if len(copies) <= MOST_COPIES:
        return
    copies.sort(key=lambda one: one.stat().st_mtime)
    for one in copies[:len(copies) - MOST_COPIES + MOST_COPIES // 10]:
        try:
            one.unlink()
        except OSError:
            pass


# -- reading ----------------------------------------------------------------------------------------
def extract(path: Any) -> Dict[str, Any]:
    """A file's reading - from its copy when there is one. Its "text" (a
    PDF's pages, each under "--- page N ---"), its "pages" for a PDF, whether
    any of it was read off a picture ("ocr"), and whether it came from the
    copy ("from_copy")."""
    path = Path(str(path))
    sha = sha256_of(path)
    kept = _kept(sha)
    if kept is not None:
        return dict(kept, from_copy=True, sha256=sha)
    kind = kind_of(path)
    said: Dict[str, Any] = {"version": VERSION, "kind": kind, "text": "", "ocr": False}
    if kind == "picture":
        said["text"] = ocr_file(path)
        said["ocr"] = True
    elif kind == "pdf":
        pages, read_off = _pdf_pages(path)
        said["pages"] = pages
        said["ocr"] = bool(read_off)
        said["ocr_pages"] = read_off
        said["text"] = "\n\n".join(f"--- page {n} ---\n{text}"
                                   for n, text in enumerate(pages, 1))
    elif kind == "xlsx":
        said["text"] = _workbook_text(path)
    else:
        said["text"] = _office_text(path, kind)
    # WORDS READ OFF A PICTURE: a password or a code in them hidden,
    # by keywords and shapes (common/credentials_guard.py) - and the picture
    # itself then never given to a model.
    if said.get("ocr") and _hides():
        from ..common.credentials_guard import hide_credentials
        hidden = hide_credentials(said["text"])
        if hidden != said["text"]:
            said["text"] = hidden
            if said.get("pages"):
                said["pages"] = hide_credentials(said["pages"])
            said["credentials"] = True
    said["made_at"] = time.time()
    said["name"] = path.name
    _keep(sha, said)
    return dict(said, from_copy=False, sha256=sha)


def _hides() -> bool:
    try:
        return bool(_hiding())
    except Exception:                              # pragma: no cover - hidden, then
        return True


def _pdf_pages(path: Path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages: List[str] = []
    read_off: List[int] = []
    for n, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:                       # a page pypdf cannot read
            text = ""
        if len(text.strip()) < SCANNED_BELOW and len(read_off) < MOST_OCR_PAGES:
            seen = _pdf_page_ocr(path, n)
            if seen.strip():
                text = seen
                read_off.append(n + 1)
        pages.append(text.strip())
    return pages, read_off


def _pdf_page_ocr(path: Path, index: int) -> str:
    """A page's words read off its picture - a scan has no text of its own."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return ""
    try:
        document = pdfium.PdfDocument(str(path))
        picture = document[index].render(scale=2).to_pil()
        document.close()
    except Exception:
        log.debug("a PDF page could not be drawn", exc_info=True)
        return ""
    return ocr_picture(picture)


def _office_text(path: Path, kind: str) -> str:
    from ..listener.providers.filesystem import _document_text
    try:
        return _document_text(str(path), "." + kind, 2_000_000)
    except Exception as exc:
        return f"(could not be read: {exc})"


def _workbook_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "(reading a workbook needs openpyxl)"
    book = load_workbook(str(path), read_only=True, data_only=True)
    lines: List[str] = []
    cells = 0
    for sheet in book.worksheets:
        lines.append(f"--- sheet {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            if cells >= MOST_CELLS:
                lines.append("(the rest left out)")
                break
            values = ["" if value is None else str(value) for value in row]
            cells += len(values)
            if any(values):
                lines.append("\t".join(values).rstrip())
    book.close()
    return "\n".join(lines)


# -- pictures ----------------------------------------------------------------------------------------
def picture_for_model(path: Any) -> Optional[Dict[str, str]]:
    """The picture, small enough for a model to be given: its media type
    and its bytes in base64 - or None when it cannot be opened."""
    try:
        from PIL import Image
        with Image.open(str(path)) as image:
            image.load()
            picture = image.convert("RGB") if image.mode not in ("RGB", "L") else image.copy()
    except Exception:
        return None
    picture.thumbnail((PICTURE_SIDE, PICTURE_SIDE))
    buffer = io.BytesIO()
    picture.save(buffer, format="PNG", optimize=True)
    data = buffer.getvalue()
    if len(data) > 1_500_000:                  # a photograph: smaller as JPEG
        buffer = io.BytesIO()
        picture.convert("RGB").save(buffer, format="JPEG", quality=85)
        return {"media_type": "image/jpeg",
                "data": base64.b64encode(buffer.getvalue()).decode("ascii")}
    return {"media_type": "image/png", "data": base64.b64encode(data).decode("ascii")}


def ocr_file(path: Any) -> str:
    try:
        from PIL import Image
        with Image.open(str(path)) as image:
            image.load()
            return ocr_picture(image.copy())
    except Exception:
        log.debug("a picture could not be opened for its words", exc_info=True)
        return ""


def ocr_picture(picture: Any) -> str:
    """The words in a picture, read on this computer by its own OCR - in each
    of the languages it has (platform_layer/ocr.py); "" when it has none."""
    from ..platform_layer.ocr import current
    ocr = current()
    if not ocr.available():
        return ""
    picture.thumbnail((OCR_SIDE, OCR_SIDE))
    handle, name = tempfile.mkstemp(suffix=".png", prefix="pa-ocr-")
    os.close(handle)
    try:
        picture.convert("RGB").save(name, format="PNG")
        return ocr.read(name)
    except Exception:
        log.debug("the OCR failed", exc_info=True)
        return ""
    finally:
        try:
            os.unlink(name)
        except OSError:
            pass


#: Pictures waiting to be read, one at a time, behind whatever took them.
_waiting: "list[str]" = []
_worker: Optional[threading.Thread] = None
#: The most waiting: a burst of pictures is not a backlog of minutes.
MOST_WAITING = 200


def read_later(path: Any) -> None:
    """A picture read now in the background - a screenshot as it is taken -
    so its copy is there before anything asks for it."""
    global _worker
    if _folder is None:
        return
    with _lock:
        if len(_waiting) >= MOST_WAITING:
            return
        _waiting.append(str(path))
        if _worker is not None and _worker.is_alive():
            return
        _worker = threading.Thread(target=_read_waiting, name="documents", daemon=True)
        _worker.start()


def _read_waiting() -> None:
    while True:
        with _lock:
            if not _waiting:
                return
            path = _waiting.pop(0)
        try:
            if Path(path).exists():
                extract(path)
        except Exception:                          # pragma: no cover - read when asked
            log.debug("a picture could not be read in the background", exc_info=True)


def read_for_model(path: Any, from_line: int = 1, most: int = 40000) -> Dict[str, Any]:
    """What a read_file tool gives a model of a picture or a document: the
    text (from the copy when there is one), from a line on; and of a picture
    the picture itself, under "_pictures" - the tools' answer is made into
    a picture block from it for a model that can see one, and it is left
    out for one that cannot (models/vision.py)."""
    said = extract(path)
    lines = str(said.get("text") or "").splitlines()
    start = max(0, int(from_line or 1) - 1)
    body = "\n".join(lines[start:])
    out: Dict[str, Any] = {
        "path": str(path), "kind": said["kind"], "from_line": start + 1,
        "text": body[:most], "cut": len(body) > most, "lines": len(lines),
        "read_off_pictures": bool(said.get("ocr")), "from_copy": bool(said.get("from_copy"))}
    if said["kind"] == "pdf":
        out["pages"] = len(said.get("pages") or [])
    if said["kind"] == "picture":
        if not out["text"]:
            out["text"] = "(no words could be read in it)"
        if said.get("credentials") and _hides():
            # Its words with the password or code hidden; the picture,
            # which would show them, not given.
            out["picture_held_back"] = "it shows a password or a code"
        else:
            picture = picture_for_model(path)
            if picture is not None:
                out["_pictures"] = [picture]
    return out


__all__ = ["configure", "extract", "copy_of", "kind_of", "needs_reading", "read_for_model",
           "picture_for_model", "ocr_picture", "ocr_file", "PICTURES", "DOCUMENTS"]
