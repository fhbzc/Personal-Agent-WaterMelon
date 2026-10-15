"""Reading what is not text, for every model (runtime/documents.py,
models/vision.py): a picture's words read on this computer and the picture
itself for a model that sees one; a PDF's pages, a scanned one read off its
picture; Office files' text; each read once and kept as a copy beside
nothing of the user's; and a model that reads text only given no picture."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from PIL import Image, ImageDraw, ImageFont

from personalized_agent.runtime import documents
from personalized_agent.models import vision


def _picture(path: Path, lines=("Quantum error correction", "Coherence doubled")) -> Path:
    canvas = Image.new("RGB", (900, 80 + 90 * len(lines)), "white")
    draw = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default()
    for n, line in enumerate(lines):
        draw.text((20, 30 + 90 * n), line, fill="black", font=font)
    canvas.save(path)
    return path


ON_WINDOWS = sys.platform == "win32"


class _Folder(unittest.TestCase):

    def setUp(self) -> None:
        self.work = tempfile.TemporaryDirectory()
        self.root = Path(self.work.name)
        self.copies = self.root / "extractions"
        documents.configure(self.copies)

    def tearDown(self) -> None:
        documents.configure(None)
        self.work.cleanup()


class PicturesTests(_Folder):

    @unittest.skipUnless(ON_WINDOWS, "Windows' own OCR")
    def test_a_picture_is_its_words_and_itself(self) -> None:
        shot = _picture(self.root / "slide.png")
        said = documents.read_for_model(shot)
        self.assertEqual(said["kind"], "picture")
        self.assertIn("Quantum", said["text"])
        self.assertTrue(said["read_off_pictures"])
        self.assertFalse(said["from_copy"])
        self.assertEqual(said["_pictures"][0]["media_type"], "image/png")
        # Read again: from the copy, nothing read twice.
        with mock.patch.object(documents, "ocr_picture",
                               side_effect=AssertionError("read again")):
            again = documents.read_for_model(shot)
        self.assertTrue(again["from_copy"])
        self.assertIn("Quantum", again["text"])
        # The copy is the application's, named by the file's bytes - never
        # beside the file.
        sha = documents.sha256_of(shot)
        self.assertTrue((self.copies / sha[:2] / f"{sha}.json").exists())
        self.assertEqual(sorted(p.name for p in self.root.iterdir()),
                         ["extractions", "slide.png"])

    @unittest.skipUnless(ON_WINDOWS, "Windows' own OCR")
    def test_a_changed_file_is_read_anew(self) -> None:
        shot = _picture(self.root / "slide.png")
        documents.read_for_model(shot)
        _picture(shot, lines=("Something else entirely",))
        said = documents.read_for_model(shot)
        self.assertFalse(said["from_copy"])
        self.assertIn("Something", said["text"])

    def test_a_big_picture_is_made_small_for_a_model(self) -> None:
        big = self.root / "big.png"
        Image.new("RGB", (5000, 3000), "white").save(big)
        picture = documents.picture_for_model(big)
        import base64
        import io
        with Image.open(io.BytesIO(base64.b64decode(picture["data"]))) as small:
            self.assertLessEqual(max(small.size), documents.PICTURE_SIDE)

    def test_the_kinds(self) -> None:
        self.assertEqual(documents.kind_of("a/B.PNG"), "picture")
        self.assertEqual(documents.kind_of("a/b.pdf"), "pdf")
        self.assertEqual(documents.kind_of("a/b.xlsm"), "xlsx")
        self.assertFalse(documents.needs_reading("a/b.py"))


class DocumentsTests(_Folder):

    @unittest.skipUnless(ON_WINDOWS, "Windows' own OCR")
    def test_a_scanned_pdf_page_is_read_off_its_picture(self) -> None:
        scan = _picture(self.root / "scan.png", lines=("Referee report two",))
        pdf = self.root / "scan.pdf"
        with Image.open(scan) as image:
            image.convert("RGB").save(pdf, "PDF")
        said = documents.read_for_model(pdf)
        self.assertEqual(said["kind"], "pdf")
        self.assertEqual(said["pages"], 1)
        self.assertIn("--- page 1 ---", said["text"])
        self.assertIn("Referee", said["text"])
        self.assertTrue(said["read_off_pictures"])

    def test_a_workbook_is_its_cells(self) -> None:
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "Budget"
        book.active.append(["item", "cost"])
        book.active.append(["chips", 120])
        path = self.root / "budget.xlsx"
        book.save(path)
        said = documents.read_for_model(path)
        self.assertIn("--- sheet Budget ---", said["text"])
        self.assertIn("chips\t120", said["text"])

    def test_a_deck_is_its_text(self) -> None:
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Results"
        path = self.root / "talk.pptx"
        deck.save(path)
        self.assertIn("Results", documents.read_for_model(path)["text"])


class WhoSeesPicturesTests(unittest.TestCase):

    def test_by_name_and_as_the_user_says(self) -> None:
        self.assertTrue(vision.sees_images("anthropic", "claude-opus-5-5"))
        self.assertTrue(vision.sees_images("codex", "default"))
        self.assertTrue(vision.sees_images("openai", "gpt-4o"))
        self.assertTrue(vision.sees_images("openrouter", "qwen/qwen2.5-vl-72b-instruct"))
        self.assertFalse(vision.sees_images("deepseek", "deepseek-chat"))
        self.assertFalse(vision.sees_images("openai", "gpt-4o-audio-preview"))

        class _Config(dict):
            def get(self, key, default=None):
                return super().get(key, default)

        said = _Config(model_sees_images={"deepseek/deepseek-chat": True,
                                          "openai/gpt-4o": False})
        self.assertTrue(vision.sees_images("deepseek", "deepseek-chat", said))
        self.assertFalse(vision.sees_images("openai", "gpt-4o", said))

    def test_a_tools_answer_with_a_picture(self) -> None:
        said = {"path": "x.png", "text": "words", "_pictures": [
            {"media_type": "image/png", "data": "AAAA"}]}
        content = vision.tool_content(said)
        self.assertEqual(content[1], {"type": "image", "source": {
            "type": "base64", "media_type": "image/png", "data": "AAAA"}})
        self.assertNotIn("_pictures", json.loads(content[0]["text"]))
        self.assertEqual(vision.tool_content({"a": 1}), '{"a": 1}')
        # For a model that reads text only: the picture said, its words kept.
        message = [{"type": "tool_result", "tool_use_id": "t", "content": content}]
        blind = vision.without_pictures(message)
        self.assertEqual(blind[0]["content"][1]["type"], "text")
        self.assertIn("words", blind[0]["content"][0]["text"])
        self.assertEqual(vision.pictures_in([{"role": "user", "content": message}]),
                         [{"media_type": "image/png", "data": "AAAA"}])
        # In a transcript, never its bytes.
        from personalized_agent.models.providers.local_tools import tool_results
        rendered = tool_results(message)
        self.assertNotIn("AAAA", rendered)
        self.assertIn("[a picture]", rendered)


class TheToolsTests(_Folder):

    def test_read_file_gives_a_picture_to_a_run(self) -> None:
        from personalized_agent.agent.work.tools import WorkTools
        shot = _picture(self.root / "fig.png")
        tools = WorkTools.__new__(WorkTools)
        tools._readable = lambda path: Path(path)
        tools._the_agents_own = lambda where: True
        with mock.patch.object(documents, "ocr_picture", return_value="Figure 3: results"):
            said = tools.read_file({"path": str(shot)})
        self.assertEqual(said["text"], "Figure 3: results")
        self.assertTrue(said["_pictures"])

    def test_the_codex_call_gets_the_pictures_as_files_and_forgets_them(self) -> None:
        from personalized_agent.models.base import ModelRequest
        from personalized_agent.models.providers.codex_provider import CodexProvider
        provider = CodexProvider.__new__(CodexProvider)
        request = ModelRequest(messages=[{"role": "user", "content": [
            {"type": "tool_result", "tool_use_id": "t", "content": vision.tool_content(
                {"text": "words", "_pictures": [{"media_type": "image/png",
                                                 "data": "iVBORw0KGgo="}]})}]}])
        paths = provider._pictures(request)
        self.assertEqual(len(paths), 1)
        self.assertTrue(Path(paths[0]).exists())
        prompt = provider._prompt(request)
        self.assertNotIn("iVBORw0KGgo=", prompt)
        self.assertIn("attached to this message", prompt)
        provider._forget_pictures(paths)
        self.assertFalse(Path(paths[0]).exists())


class EachSystemsOcrTests(unittest.TestCase):

    def test_the_reading_most_of_its_own_script_is_kept(self) -> None:
        from personalized_agent.platform_layer.ocr import best_reading
        chinese = {"en": "lI ii I1 l", "zh": "\u91cf\u5b50\u7ea0\u9519 2026"}
        self.assertEqual(best_reading(chinese, ["en", "zh"]), chinese["zh"])
        english = {"en": "Quantum error correction", "zh": "Quantum error correctlon"}
        self.assertEqual(best_reading(english, ["en", "zh"]), english["en"])
        self.assertEqual(best_reading({}, []), "")

    def test_windows_reads_in_each_language_it_has(self) -> None:
        from personalized_agent.platform_layer import ocr
        windows = ocr.WindowsOcr()
        windows._known, windows._known_at = ["en", "zh"], 1e18
        windows._tags = {"en": "en-US", "zh": "zh-Hans-CN"}
        asked = {}

        def run(script, env=None, timeout=120):
            asked.update(env or {})
            out = ("@@LANG en-US\nlI ii\n@@LANG zh-Hans-CN\n\u91cf\u5b50\u7ea0\u9519\n")
            return mock.Mock(returncode=0, stdout=out.encode("utf-8"))

        with mock.patch.object(ocr, "_powershell", side_effect=run):
            self.assertEqual(windows.read("x.png"), "\u91cf\u5b50\u7ea0\u9519")
        self.assertEqual(asked["PA_OCR_LANGS"], "en-US,zh-Hans-CN")

    def test_windows_adds_languages_only_with_the_administrators_yes(self) -> None:
        from personalized_agent.platform_layer import ocr
        windows = ocr.WindowsOcr()
        with mock.patch.object(ocr.subprocess, "Popen") as launch:
            said = windows.add(["zh", "ja"])
        self.assertTrue(said["started"])
        command = " ".join(launch.call_args.args[0])
        self.assertIn("-Verb RunAs", command)
        self.assertIn("Language.OCR~~~zh-CN~0.0.1.0", command)
        self.assertIn("Language.OCR~~~ja-JP~0.0.1.0", command)

    def test_linux_says_what_the_system_must_add(self) -> None:
        from personalized_agent.platform_layer import ocr
        linux = ocr.LinuxOcr()
        with mock.patch.object(ocr.shutil, "which", return_value=None), \
                mock.patch.object(ocr.LinuxOcr, "_rapid", return_value=True):
            self.assertEqual(linux.installed(), ["en", "zh"])
            said = linux.add(["ja"])
            status = linux.status()
        self.assertFalse(said["started"])
        self.assertIn("tesseract-ocr-jpn", said["how"])
        self.assertEqual([one["code"] for one in status["languages"] if one["installed"]],
                         ["en", "zh"])

    def test_no_ocr_no_words(self) -> None:
        from personalized_agent.platform_layer import ocr
        blind = ocr.Ocr()
        with mock.patch.object(ocr, "current", return_value=blind):
            self.assertEqual(documents.ocr_picture(Image.new("RGB", (10, 10))), "")
        self.assertFalse(blind.status()["available"])


class TheManagerTests(unittest.TestCase):
    """A model that reads text only is given no picture - each said where
    it was; one that sees them is given them."""

    def test_text_only_or_pictures(self) -> None:
        sys.path.insert(0, str(Path(__file__).parent))
        from support import TemporaryAgent
        from personalized_agent.models.base import ModelRequest, ModelResponse
        agent = TemporaryAgent()
        try:
            models = agent.app.models
            resolved = models.resolve("primary")
            given = []
            provider = models.provider(resolved.provider_id)
            content = vision.tool_content({"text": "words", "_pictures": [
                {"media_type": "image/png", "data": "AAAA"}]})
            request = ModelRequest(messages=[{"role": "user", "content": [
                {"type": "tool_result", "tool_use_id": "t", "content": content}]}])
            with mock.patch.object(provider, "invoke", side_effect=lambda req, model: (
                    given.append(req) or ModelResponse(text="ok"))):
                models.invoke(request, role="primary")
                agent.app.config.set("model_sees_images", {
                    f"{resolved.provider_id}/{resolved.model_id}": True}, internal=True)
                models.invoke(request, role="primary")
            self.assertEqual(vision.pictures_in(given[0].messages), [])
            self.assertIn("reads text only", json.dumps(given[0].messages))
            self.assertEqual(len(vision.pictures_in(given[1].messages)), 1)
            self.assertEqual(len(vision.pictures_in(request.messages)), 1,
                             "the caller's request is not changed")
        finally:
            agent.cleanup()


if __name__ == "__main__":
    unittest.main()
